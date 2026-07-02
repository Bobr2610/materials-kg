"""Portable hypothesis reports for expert review outside the application."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from kg_engine.domain.models import HypothesisGenerationResult


def render_hypothesis_report(
    result: HypothesisGenerationResult,
    report_format: str,
) -> bytes:
    """Render a versioned report in a supported portable format."""
    renderers = {
        "markdown": _render_markdown,
        "xlsx": _render_xlsx,
        "docx": _render_docx,
        "pdf": _render_pdf,
    }
    try:
        renderer = renderers[report_format]
    except KeyError as exc:
        raise ValueError(f"Unsupported report format: {report_format}") from exc
    return renderer(result)


def _render_markdown(result: HypothesisGenerationResult) -> bytes:
    lines = [
        f"# Hypothesis report: {result.target_kpi}",
        "",
        f"Generation engine: `{result.generation_engine}`",
        "",
    ]
    for hypothesis in result.hypotheses:
        lines.extend(
            [
                f"## {hypothesis.rank or '-'} — {hypothesis.statement}",
                "",
                f"- Score: {hypothesis.score.final_score:.3f}",
                f"- Mechanism: {hypothesis.mechanism or 'Not specified'}",
                f"- Rationale: {hypothesis.rationale}",
                f"- Test plan: {hypothesis.test_plan}",
                "- Evidence: "
                + ", ".join(
                    [
                        *hypothesis.supporting_evidence_ids,
                        *hypothesis.supporting_observation_ids,
                        *hypothesis.supporting_text_unit_ids,
                    ]
                ),
                "",
            ]
        )
    return "\n".join(lines).encode("utf-8")


def _render_xlsx(result: HypothesisGenerationResult) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hypotheses"
    sheet.append(
        [
            "Rank",
            "ID",
            "Statement",
            "Mechanism",
            "Final score",
            "Novelty",
            "Risk",
            "Value",
            "Evidence strength",
            "Evidence IDs",
            "Test plan",
        ]
    )
    for item in result.hypotheses:
        sheet.append(
            [
                item.rank,
                item.id,
                item.statement,
                item.mechanism,
                item.score.final_score,
                item.score.novelty,
                item.score.risk,
                item.score.value,
                item.score.evidence_strength,
                ";".join(
                    [
                        *item.supporting_evidence_ids,
                        *item.supporting_observation_ids,
                        *item.supporting_text_unit_ids,
                    ]
                ),
                item.test_plan,
            ]
        )
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _render_docx(result: HypothesisGenerationResult) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(f"Hypothesis report: {result.target_kpi}", level=0)
    document.add_paragraph(f"Generation engine: {result.generation_engine}")
    for item in result.hypotheses:
        document.add_heading(f"{item.rank or '-'} — {item.statement}", level=1)
        document.add_paragraph(f"Final score: {item.score.final_score:.3f}")
        document.add_heading("Mechanism", level=2)
        document.add_paragraph(item.mechanism or "Not specified")
        document.add_heading("Rationale", level=2)
        document.add_paragraph(item.rationale)
        document.add_heading("Verification", level=2)
        document.add_paragraph(item.test_plan)
        document.add_heading("Provenance IDs", level=2)
        evidence_ids = [
            *item.supporting_evidence_ids,
            *item.supporting_observation_ids,
            *item.supporting_text_unit_ids,
        ]
        document.add_paragraph(", ".join(evidence_ids) or "No evidence IDs")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _render_pdf(result: HypothesisGenerationResult) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas

    output = BytesIO()
    canvas = Canvas(output, pagesize=A4)
    font_name = _register_pdf_font(pdfmetrics, TTFont)
    _width, height = A4
    y = height - 48

    def write(text: str, *, size: int = 10) -> None:
        nonlocal y
        if y < 48:
            canvas.showPage()
            y = height - 48
        canvas.setFont(font_name, size)
        for line in _wrap_text(text, 95):
            canvas.drawString(48, y, line)
            y -= size + 4

    write(f"Hypothesis report: {result.target_kpi}", size=15)
    write(f"Generation engine: {result.generation_engine}")
    for item in result.hypotheses:
        y -= 8
        write(f"{item.rank or '-'} — {item.statement}", size=12)
        write(f"Score: {item.score.final_score:.3f}")
        write(f"Mechanism: {item.mechanism or 'Not specified'}")
        write(f"Rationale: {item.rationale}")
        write(f"Verification: {item.test_plan}")
        ids = [
            *item.supporting_evidence_ids,
            *item.supporting_observation_ids,
            *item.supporting_text_unit_ids,
        ]
        write(f"Provenance IDs: {', '.join(ids) or 'none'}")
    canvas.save()
    return output.getvalue()


def _register_pdf_font(pdfmetrics: Any, font_type: Any) -> str:
    candidates = (
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    for path in candidates:
        if path.exists():
            name = "MaterialsReportUnicode"
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(font_type(name, str(path)))
            return name
    return "Helvetica"


def _wrap_text(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current: list[str] = []
    for word in words:
        if current and len(" ".join([*current, word])) > width:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    return lines or [""]

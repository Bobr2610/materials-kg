"""Safe format registry for text, PDF, DOCX, XLSX, and raster images."""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import ClassVar
from zipfile import BadZipFile
from zipfile import ZipFile

from kg_engine.domain.ingestion import IngestionWarning
from kg_engine.domain.ingestion import ParsedFragment
from kg_engine.domain.ingestion import ParsedSource
from kg_engine.domain.ingestion import ParsedTable

MAX_SOURCE_BYTES = 50 * 1024 * 1024


class UnsupportedDocumentError(ValueError):
    """Raised when a file type is outside the reader allowlist."""


class InvalidDocumentSignatureError(ValueError):
    """Raised when bytes do not match the allowlisted filename format."""


class DocumentReaderRegistry:
    """Dispatch bytes to a deterministic, format-specific reader."""

    MEDIA_TYPES: ClassVar[dict[str, str]] = {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".tif": "image/tiff",
        ".tiff": "image/tiff",
    }

    def is_supported(self, name: str) -> bool:
        """Return whether a filename has an allowlisted suffix."""
        return Path(name).suffix.lower() in self.MEDIA_TYPES

    def read_path(self, path: str | Path) -> ParsedSource:
        """Read one local file after validating its size."""
        source_path = Path(path)
        if source_path.stat().st_size > MAX_SOURCE_BYTES:
            raise ValueError(f"Source exceeds {MAX_SOURCE_BYTES} bytes")
        return self.read_bytes(source_path.name, source_path.read_bytes())

    def read_bytes(self, name: str, content: bytes) -> ParsedSource:
        """Parse bytes without trusting a caller-provided MIME type."""
        suffix = Path(name).suffix.lower()
        if suffix not in self.MEDIA_TYPES:
            raise UnsupportedDocumentError(f"Unsupported file type: {suffix or '<none>'}")
        if len(content) > MAX_SOURCE_BYTES:
            raise ValueError(f"Source exceeds {MAX_SOURCE_BYTES} bytes")
        self._validate_signature(suffix, content)
        base = ParsedSource(
            name=Path(name).name,
            media_type=self.MEDIA_TYPES[suffix],
            checksum=sha256(content).hexdigest(),
        )
        if not content:
            base.warnings.append(
                IngestionWarning(code="empty_source", message="Source file is empty")
            )
            return base
        if suffix in {".txt", ".md"}:
            return self._read_text(base, content)
        if suffix == ".docx":
            return self._read_docx(base, content)
        if suffix == ".xlsx":
            return self._read_xlsx(base, content)
        if suffix == ".pdf":
            return self._read_pdf(base, content)
        return self._read_image(base, content)

    @staticmethod
    def _validate_signature(suffix: str, content: bytes) -> None:
        valid = True
        if suffix == ".pdf":
            valid = content.startswith(b"%PDF-")
        elif suffix in {".docx", ".xlsx"}:
            expected = "word/document.xml" if suffix == ".docx" else "xl/workbook.xml"
            try:
                with ZipFile(BytesIO(content)) as archive:
                    valid = expected in archive.namelist()
            except BadZipFile:
                valid = False
        elif suffix == ".png":
            valid = content.startswith(b"\x89PNG\r\n\x1a\n")
        elif suffix in {".jpg", ".jpeg"}:
            valid = content.startswith(b"\xff\xd8\xff")
        elif suffix in {".tif", ".tiff"}:
            valid = content.startswith((b"II*\x00", b"MM\x00*"))
        if not valid:
            raise InvalidDocumentSignatureError(
                f"File signature does not match {suffix} format"
            )

    @staticmethod
    def _read_text(source: ParsedSource, content: bytes) -> ParsedSource:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
            source.warnings.append(
                IngestionWarning(
                    code="encoding_fallback",
                    message="Decoded source using latin-1 fallback",
                )
            )
        source.fragments = [
            ParsedFragment(content=line, locator={"line": index})
            for index, line in enumerate(text.splitlines(), start=1)
            if line.strip()
        ]
        return source

    @staticmethod
    def _read_docx(source: ParsedSource, content: bytes) -> ParsedSource:
        from docx import Document

        document = Document(BytesIO(content))
        source.fragments = [
            ParsedFragment(content=item.text, locator={"paragraph": index})
            for index, item in enumerate(document.paragraphs, start=1)
            if item.text.strip()
        ]
        for index, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            source.tables.append(
                ParsedTable(
                    headers=rows[0] if rows else [],
                    rows=rows,
                    locator={"table": index},
                )
            )
        return source

    @staticmethod
    def _read_xlsx(source: ParsedSource, content: bytes) -> ParsedSource:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
        for sheet in workbook.worksheets:
            rows = [
                ["" if value is None else str(value) for value in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            while rows and not any(rows[-1]):
                rows.pop()
            if not rows:
                continue
            width = max(len(row) for row in rows)
            table_range = f"A1:{get_column_letter(width)}{len(rows)}"
            source.tables.append(
                ParsedTable(
                    headers=rows[0],
                    rows=rows,
                    locator={"sheet": sheet.title, "range": table_range},
                )
            )
            source.fragments.extend(
                ParsedFragment(
                    kind="table_row",
                    content=" | ".join(row),
                    locator={"sheet": sheet.title, "row": row_number},
                )
                for row_number, row in enumerate(rows, start=1)
            )
        return source

    @staticmethod
    def _read_pdf(source: ParsedSource, content: bytes) -> ParsedSource:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            source.warnings.append(
                IngestionWarning(code="encrypted_pdf", message="PDF is encrypted")
            )
            return source
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                source.fragments.append(
                    ParsedFragment(content=text, locator={"page": page_number})
                )
            else:
                source.warnings.append(
                    IngestionWarning(
                        code="ocr_recommended",
                        message="No text layer found on PDF page",
                        locator={"page": page_number},
                    )
                )
        return source

    @staticmethod
    def _read_image(source: ParsedSource, content: bytes) -> ParsedSource:
        from PIL import Image
        import pytesseract

        image = Image.open(BytesIO(content))
        text = pytesseract.image_to_string(image).strip()
        if text:
            source.fragments.append(
                ParsedFragment(kind="ocr", content=text, locator={"image": 1})
            )
        else:
            source.warnings.append(
                IngestionWarning(code="empty_ocr", message="OCR returned no text")
            )
        return source
